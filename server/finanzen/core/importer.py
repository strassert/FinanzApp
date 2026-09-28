"""Statement import (CSV, XLSX, XLS) for accounts without API.

Two steps: `preview()` detects the header and guesses a column mapping,
`run_import()` imports with the mapping the user confirmed (remembered per
account). Transactions without a bank reference get a fingerprint identity,
so re-importing an overlapping file only adds what is new.
"""

from __future__ import annotations

import csv
import io
import json
import re
import sqlite3
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from typing import Optional

from ..money import to_minor
from . import depot
from .store import NewTx, store_balance, upsert_transactions

MAX_ROWS = 20000


class ImportError_(Exception):
    """User-facing import error (German message)."""


# --- file reading ---------------------------------------------------------------------

def detect_kind(data: bytes) -> str:
    if data[:4] == b"%PDF":
        return "pdf"
    if data[:4] == b"PK\x03\x04":
        return "xlsx"
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "xls"
    return "csv"


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return f"{value:.2f}"
    return str(value).strip()


def read_rows(data: bytes) -> list[list[str]]:
    kind = detect_kind(data)
    if kind == "pdf":
        raise ImportError_("PDF-Dateien lassen sich nicht zuverlässig lesen. Bitte exportiere die Umsätze "
                           "im Online-Banking als CSV oder Excel.")
    if kind == "xlsx":
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise ImportError_("Die Excel-Datei konnte nicht gelesen werden.") from exc
        ws = wb.worksheets[0]
        rows = [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True, max_row=MAX_ROWS)]
        wb.close()
        return _trim(rows)
    if kind == "xls":
        import xlrd
        try:
            book = xlrd.open_workbook(file_contents=data)
        except Exception as exc:
            raise ImportError_("Die Excel-Datei (.xls) konnte nicht gelesen werden.") from exc
        sheet = book.sheet_by_index(0)
        rows = []
        for r in range(min(sheet.nrows, MAX_ROWS)):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(_cell(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)))
                elif cell.ctype == xlrd.XL_CELL_NUMBER:
                    row.append(f"{cell.value:.2f}")
                else:
                    row.append(str(cell.value).strip())
            rows.append(row)
        return _trim(rows)
    text = _decode(data)
    sample = "\n".join(text.splitlines()[:50])
    delimiter = max([";", ",", "\t", "|"], key=lambda d: sample.count(d))
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))[:MAX_ROWS]
    return _trim([[c.strip() for c in r] for r in rows])


def _trim(rows: list[list[str]]) -> list[list[str]]:
    return [r for r in rows if any(c for c in r)]


# --- value parsing -----------------------------------------------------------------------

_CURRENCY_SIGNS = {"€": "EUR", "$": "USD", "£": "GBP", "CHF": "CHF", "EUR": "EUR", "USD": "USD", "GBP": "GBP"}


def parse_amount(text: str) -> Optional[Decimal]:
    """1.234,56 | 1,234.56 | 12,50- | (3.20) | −12,50 | € 12,50 | 12.50 EUR | +5"""
    if text is None:
        return None
    s = str(text).strip().replace("−", "-").replace(" ", "").replace(" ", "").replace("'", "")
    if not s:
        return None
    s = re.sub(r"(?i)EUR|USD|GBP|CHF|[€$£]", "", s)
    negative = False
    if s.startswith("(") and s.endswith(")"):
        negative, s = True, s[1:-1]
    if s.endswith("-"):
        negative, s = True, s[:-1]
    if s.startswith("-"):
        negative, s = True, s[1:]
    s = s.lstrip("+")
    if not re.fullmatch(r"[\d.,]+", s or "x"):
        return None
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        # one comma = German decimal separator (12,50); several = English thousands (1,234,567)
        s = s.replace(",", "") if s.count(",") > 1 else s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.rpartition(".")[2]) == 3):
        s = s.replace(".", "")          # German thousands: 1.234 / 1.234.567
    try:
        value = Decimal(s)
    except InvalidOperation:
        return None
    return -value if negative else value


def detect_currency(text: str) -> Optional[str]:
    for sign, code in _CURRENCY_SIGNS.items():
        if sign in str(text):
            return code
    return None


_DATE_FORMATS = ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%m/%d/%Y")


def parse_date(text: str, fmt: Optional[str] = None) -> Optional[date]:
    s = str(text or "").strip()
    if not s:
        return None
    s = re.split(r"[ T]", s, maxsplit=1)[0] if re.match(r"\d{4}-\d{2}-\d{2}[ T]", s) else s.split(" ")[0]
    for f in ((fmt,) if fmt else _DATE_FORMATS):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            continue
    return None


# --- mapping -----------------------------------------------------------------------------

@dataclass
class Mapping:
    header_row: int
    date: int
    amount: Optional[int] = None
    debit: Optional[int] = None
    credit: Optional[int] = None
    text: tuple[int, ...] = ()
    counterparty: Optional[int] = None
    currency: Optional[int] = None
    balance: Optional[int] = None
    invert: bool = False
    date_format: Optional[str] = None
    skip_column: Optional[int] = None
    skip_values: tuple[str, ...] = ()

    def to_json(self) -> dict:
        return {k: (list(v) if isinstance(v, tuple) else v) for k, v in self.__dict__.items()}

    @classmethod
    def from_json(cls, data: dict) -> "Mapping":
        known = {k: v for k, v in data.items() if k in cls.__dataclass_fields__}
        for key in ("text", "skip_values"):
            if key in known:
                known[key] = tuple(known[key] or ())
        return cls(**known)


TEMPLATES = [
    # (name, required headers, mapping by header name)
    ("PayPal (Deutsch)", {"Datum", "Brutto", "Währung", "Name"},
     {"date": "Datum", "amount": "Brutto", "currency": "Währung", "counterparty": "Name",
      "text": ["Typ", "Artikelbezeichnung"], "balance": "Guthaben", "date_format": "%d.%m.%Y",
      "skip_column": "Status", "skip_values": ["Ausstehend", "Storniert", "Abgelehnt"]}),
    ("PayPal (English)", {"Date", "Gross", "Currency", "Name"},
     {"date": "Date", "amount": "Gross", "currency": "Currency", "counterparty": "Name",
      "text": ["Type", "Item Title"], "balance": "Balance", "date_format": "%d/%m/%Y",
      "skip_column": "Status", "skip_values": ["Pending", "Denied", "Reversed"]}),
    ("Revolut", {"Started Date", "Amount", "Currency", "Description"},
     {"date": "Completed Date", "amount": "Amount", "currency": "Currency", "text": ["Description"],
      "balance": "Balance", "skip_column": "State", "skip_values": ["PENDING", "REVERTED", "DECLINED"]}),
]

_DATE_HEADERS = ("buchungsdatum", "buchungstag", "datum", "date", "valuta", "wertstellung", "transaktionsdatum")
_AMOUNT_HEADERS = ("betrag", "amount", "brutto", "gross", "umsatz", "summe", "wert")
_DEBIT_HEADERS = ("soll", "belastung", "debit", "ausgang", "lastschrift")
_CREDIT_HEADERS = ("haben", "gutschrift", "credit", "eingang")
_TEXT_HEADERS = ("buchungstext", "umsatztext", "verwendungszweck", "beschreibung", "description", "text",
                 "details", "zahlungsreferenz", "händler", "merchant")
_PARTY_HEADERS = ("empfänger", "auftraggeber", "name", "partner", "gegenkonto", "zahlungsempfänger")
_CURRENCY_HEADERS = ("währung", "currency", "whg")
_BALANCE_HEADERS = ("saldo", "balance", "kontostand", "guthaben")


def find_header(rows: list[list[str]]) -> int:
    best, best_score = 0, -1
    for i, row in enumerate(rows[:30]):
        cells = [c.lower() for c in row if c]
        score = sum(any(h in c for h in _DATE_HEADERS + _AMOUNT_HEADERS + _TEXT_HEADERS + _DEBIT_HEADERS)
                    for c in cells)
        if len(cells) >= 2 and score > best_score:
            best, best_score = i, score
    return best


def _col_ok(rows: list[list[str]], col: int, check) -> bool:
    values = [r[col] for r in rows if col < len(r) and r[col]]
    return bool(values) and sum(1 for v in values if check(v) is not None) >= 0.8 * len(values)


def guess_mapping(rows: list[list[str]]) -> tuple[Mapping, Optional[str]]:
    header_row = find_header(rows)
    header = rows[header_row]
    body = rows[header_row + 1:header_row + 51]
    names = [h.strip() for h in header]

    for name, required, spec in TEMPLATES:
        if required <= set(names):
            idx = {n: i for i, n in enumerate(names)}
            date_col = idx.get(spec["date"], idx.get("Started Date"))
            return Mapping(
                header_row=header_row, date=date_col, amount=idx[spec["amount"]],
                currency=idx.get(spec.get("currency")), counterparty=idx.get(spec.get("counterparty")),
                text=tuple(idx[t] for t in spec.get("text", []) if t in idx),
                balance=idx.get(spec.get("balance")), date_format=spec.get("date_format"),
                skip_column=idx.get(spec.get("skip_column")), skip_values=tuple(spec.get("skip_values", ())),
            ), name

    lower = [n.lower() for n in names]

    def find(keys, check=None, exclude=()):
        for key in keys:
            for i, h in enumerate(lower):
                if key in h and i not in exclude and not any(t in h for t in ("text", "zweck")) \
                        and (check is None or _col_ok(body, i, check)):
                    return i
        return None

    date_col = find(_DATE_HEADERS, parse_date)
    if date_col is None:
        date_col = next((i for i in range(len(names)) if _col_ok(body, i, parse_date)), None)
    if date_col is None:
        raise ImportError_("Keine Datumsspalte gefunden.")
    amount = find(_AMOUNT_HEADERS, parse_amount, exclude={date_col})
    debit = credit = None
    if amount is None:
        debit = find(_DEBIT_HEADERS, parse_amount, exclude={date_col})
        credit = find(_CREDIT_HEADERS, parse_amount, exclude={date_col})
    if amount is None and (debit is None or credit is None):
        raise ImportError_("Keine Betragsspalte gefunden.")
    used = {date_col, amount, debit, credit}
    text = tuple(i for i, h in enumerate(lower) if i not in used and any(k in h for k in _TEXT_HEADERS))
    party = next((i for i, h in enumerate(lower) if i not in used and i not in text
                  and any(k in h for k in _PARTY_HEADERS)), None)
    currency = next((i for i, h in enumerate(lower) if any(k == h or k in h for k in _CURRENCY_HEADERS)), None)
    balance = next((i for i, h in enumerate(lower) if i not in used and any(k in h for k in _BALANCE_HEADERS)
                    and _col_ok(body, i, parse_amount)), None)
    if not text:
        text = tuple(i for i in range(len(names)) if i not in used and i not in (party, currency, balance)
                     and any(r[i] for r in body if i < len(r)) and not _col_ok(body, i, parse_amount))[:2]
    return Mapping(header_row, date_col, amount, debit, credit, text, party, currency, balance), None


# --- conversion --------------------------------------------------------------------------

@dataclass
class Parsed:
    txs: list[NewTx]
    skipped: int
    last_balance: Optional[tuple[date, int, str]]


def parse_rows(rows: list[list[str]], m: Mapping, default_currency: str) -> Parsed:
    out: list[NewTx] = []
    skipped = 0
    last_balance = None

    def cell(row, col):
        return row[col] if col is not None and col < len(row) else ""

    for row in rows[m.header_row + 1:]:
        if m.skip_column is not None and cell(row, m.skip_column) in m.skip_values:
            skipped += 1
            continue
        day = parse_date(cell(row, m.date), m.date_format)
        if m.amount is not None:
            amount = parse_amount(cell(row, m.amount))
        else:
            debit, credit = parse_amount(cell(row, m.debit)), parse_amount(cell(row, m.credit))
            amount = None if debit is None and credit is None else (credit or 0) - abs(debit or 0)
        if day is None or amount is None:
            skipped += 1
            continue
        if m.invert:
            amount = -amount
        currency = (cell(row, m.currency).upper() if m.currency is not None and cell(row, m.currency)
                    else detect_currency(cell(row, m.amount) if m.amount is not None else "") or default_currency)
        text = " ".join(cell(row, c) for c in m.text if cell(row, c)).strip()
        out.append(NewTx(booking_date=day, amount_minor=to_minor(amount, currency), currency=currency,
                         counterparty=cell(row, m.counterparty) or None, description=text))
        if m.balance is not None:
            bal = parse_amount(cell(row, m.balance))
            if bal is not None and (last_balance is None or day >= last_balance[0]):
                last_balance = (day, to_minor(bal, currency), currency)
    return Parsed(out, skipped, last_balance)


def _depot_account(conn: sqlite3.Connection, account_id: int) -> None:
    kind = conn.execute("SELECT kind FROM accounts WHERE id=?", (account_id,)).fetchone()[0]
    if kind != "depot":
        raise ImportError_("Das ist ein Depot-Export. Bitte ein Depot-Konto wählen "
                           "(unter „Konten“ als Art „Depot“ anlegen).")


def preview(conn: sqlite3.Connection, data: bytes, account_id: int) -> dict:
    rows = read_rows(data)
    if len(rows) < 2:
        raise ImportError_("Die Datei enthält keine Umsätze.")
    if depot.is_flatex_depot(rows):
        _depot_account(conn, account_id)
        trades, skipped = depot.parse_flatex_depot(rows)
        totals: dict[str, dict] = {}
        for t in trades:
            item = totals.setdefault(t.isin, {"isin": t.isin, "name": t.name, "quantity": Decimal(0)})
            item["quantity"] += t.quantity
        return {"kind": "depot", "template": "flatex Depotumsätze", "count": len(trades), "skipped": skipped,
                "holdings": [dict(h, quantity=str(h["quantity"])) for h in totals.values() if h["quantity"]]}
    saved = conn.execute("SELECT mapping FROM import_mappings WHERE account_id=?", (account_id,)).fetchone()
    template = None
    if saved:
        mapping = Mapping.from_json(json.loads(saved[0]))
        header = rows[mapping.header_row] if mapping.header_row < len(rows) else []
        if len(header) <= max(i for i in (mapping.date, mapping.amount or 0, mapping.debit or 0) if i is not None):
            mapping, template = guess_mapping(rows)
    else:
        mapping, template = guess_mapping(rows)
    currency = conn.execute("SELECT currency FROM accounts WHERE id=?", (account_id,)).fetchone()[0]
    parsed = parse_rows(rows, mapping, currency)
    return {
        "columns": rows[mapping.header_row],
        "sample": rows[mapping.header_row + 1:mapping.header_row + 11],
        "mapping": mapping.to_json(),
        "template": template,
        "remembered": bool(saved),
        "count": len(parsed.txs),
        "skipped": parsed.skipped,
        "parsed": [{"date": t.booking_date.isoformat(), "amount": t.amount_minor, "currency": t.currency,
                    "text": t.description, "counterparty": t.counterparty} for t in parsed.txs[:10]],
        "has_balance": parsed.last_balance is not None,
    }


def run_import(conn: sqlite3.Connection, data: bytes, account_id: int, mapping_json: dict,
               today: Optional[date] = None) -> dict:
    rows = read_rows(data)
    if depot.is_flatex_depot(rows):
        _depot_account(conn, account_id)
        trades, skipped = depot.parse_flatex_depot(rows)
        if not trades:
            raise ImportError_("Die Datei enthält keine Depotumsätze.")
        new, known = depot.import_trades(conn, account_id, trades)
        depot.revalue(conn, account_id, today or date.today())
        return {"new": new, "known": known, "skipped": skipped,
                "first": min(t.booking_date for t in trades).isoformat(),
                "last": max(t.booking_date for t in trades).isoformat(),
                "holdings": depot.holdings_summary(conn, account_id, today or date.today())}
    mapping = Mapping.from_json(mapping_json)
    currency = conn.execute("SELECT currency FROM accounts WHERE id=?", (account_id,)).fetchone()[0]
    parsed = parse_rows(rows, mapping, currency)
    if not parsed.txs:
        raise ImportError_("Mit dieser Zuordnung wurden keine Umsätze erkannt.")
    result = upsert_transactions(conn, account_id, "import", parsed.txs)
    conn.execute("INSERT OR REPLACE INTO import_mappings (account_id, mapping) VALUES (?,?)",
                 (account_id, json.dumps(mapping.to_json())))
    if parsed.last_balance:
        day, amount, cur = parsed.last_balance
        store_balance(conn, account_id, amount, cur, "IMPORT", day)
    return {"new": len(result.new_ids), "known": len(result.seen_ids) - len(result.new_ids),
            "skipped": parsed.skipped,
            "first": min(t.booking_date for t in parsed.txs).isoformat(),
            "last": max(t.booking_date for t in parsed.txs).isoformat()}
